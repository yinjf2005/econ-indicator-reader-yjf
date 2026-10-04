#!/usr/bin/env python
"""
Word 原生方程装配：LaTeX -> Word OMML (via pandoc stdout, 零临时文件)

用法：
    import sys, os
    sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "vizmath"))
    from tex2omml import add_math_run, omml_for

    p = doc.add_paragraph()
    add_math_run(p, r"\\frac{\\partial \\Delta D_t}{\\partial p_t}=\\bar\\theta\\kappa S_{t-1}>0")

命令行自检：
    python tex2omml.py --selftest
    python tex2omml.py "E=mc^2"            # 打印转换出的 m:t 词元
"""
import io
import os
import re
import subprocess
import sys

__all__ = ["omml_for", "add_math_run", "clean_tex", "pandoc_path", "STRUCT_TAGS"]

MATH_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
STRUCT_TAGS = ("m:f", "m:sSup", "m:sSub", "m:sup", "m:sub", "m:acc", "m:rad", "m:d")

_OMML_CACHE = {}
_FAILED = []


def _skill_root():
    """_shared/vizmath/scripts -> 上两级"""
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def pandoc_path():
    try:
        import pypandoc
        return pypandoc.get_pandoc_path()
    except Exception:  # 回退：搜索受管 venv
        base = os.path.expanduser("~/.workbuddy/binaries/python/envs/default")
        for cand in [
            os.path.join(base, "Scripts", "pandoc.exe"),
            os.path.join(base, "Lib", "site-packages", "pypandoc", "files", "pandoc.exe"),
        ]:
            if os.path.exists(cand):
                return cand
        raise RuntimeError(
            "未找到 pandoc。安装："
            "~/.workbuddy/binaries/python/envs/default/Scripts/pip.exe install pypandoc-binary"
        )


def clean_tex(tex):
    """把中文 \\text{...} 摘出来（OMML 缺字形会变方框），返回 (净化后的tex, 中文含义列表)。"""
    glosses = []
    def _repl(m):
        inner = m.group(1)
        if re.search(r"[\u4e00-\u9fff]", inner):
            glosses.append(inner)
            return ""
        return r"\%s{%s}" % ("mathrm", inner)  # 纯 ASCII 的 text -> mathrm
    cleaned = re.sub(r"\\text\s*\{([^{}]*)\}", _repl, tex)
    cleaned = re.sub(r"\s+", " ", cleaned).strip()
    return cleaned, glosses


def omml_for(tex):
    """返回 <m:oMath> 元素（deepcopy 前请勿跨文档复用同一节点）。失败返回 None 并记入 _FAILED。"""
    cleaned, glosses = clean_tex(tex.strip())
    key = cleaned
    if key in _OMML_CACHE:
        return _OMML_CACHE[key]

    el = None
    src = "${}$".format(key)
    try:
        from docx import Document
        proc = subprocess.run(
            [pandoc_path(), "-f", "latex", "-t", "docx", "-o", "-"],
            input=src.encode("utf-8"), capture_output=True, timeout=60,
        )
        if proc.returncode != 0 or not proc.stdout:
            raise RuntimeError(proc.stderr.decode("utf-8", "ignore")[:180])
        doc = Document(io.BytesIO(proc.stdout))
        for para in doc.paragraphs:
            node = para._p.find(MATH_NS + "oMathPara")
            if node is not None:
                inner = node.find(MATH_NS + "oMath")
                node = inner if inner is not None else node
            else:
                node = para._p.find(MATH_NS + "oMath")
            if node is None:
                node = para._p.find(".//" + MATH_NS + "oMath")
            if node is not None:
                el = node
                break
        if el is None:
            raise RuntimeError("pandoc 未产出 oMath 节点")
    except Exception as ex:
        _FAILED.append((key[:70], str(ex)[:120]))
        el = None

    _OMML_CACHE[key] = el
    return el


def add_math_run(paragraph, tex, before="", after=""):
    """往段落追加一段原生方程；失败则回退纯文本（有声回退并记录）。"""
    import copy
    from docx.oxml.ns import qn

    if before:
        paragraph.add_run(before)
    cleaned, glosses = clean_tex(tex.strip())
    el = omml_for(cleaned)
    if el is not None:
        paragraph._p.append(copy.deepcopy(el))
    else:
        paragraph.add_run("[公式未原生化: %s]" % clean_display(cleaned))
    if glosses:
        paragraph.add_run("（%s）" % "；".join(glosses))
    if after:
        paragraph.add_run(after)
    return paragraph


def clean_display(tex):
    """把 LaTeX 转成可读纯文本用于回退显示（近似，只做最常见替换）。"""
    s = tex
    s = re.sub(r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}", r"(\1)/(\2)", s)
    s = re.sub(r"\\partial\s*", "∂", s)
    s = re.sub(r"\\bar\s*\{([^{}]*)\}", r"\1‾", s)
    s = re.sub(r"\\Delta\s*", "Δ", s)
    s = re.sub(r"\\theta\s*", "θ", s)
    s = re.sub(r"\\kappa\s*", "κ", s)
    s = re.sub(r"\\phi\s*", "φ", s)
    s = re.sub(r"\\lambda\s*", "λ", s)
    s = re.sub(r"\\sigma\s*", "σ", s)
    s = re.sub(r"\\varepsilon\s*", "ε", s)
    s = re.sub(r"\\eta\s*", "η", s)
    s = re.sub(r"\\alpha\s*", "α", s)
    s = re.sub(r"\\beta\s*", "β", s)
    s = re.sub(r"_{([^{}]*)}", r"_(\1)", s)
    s = re.sub(r"\^{([^{}]*)}", r"^(\1)", s)
    s = s.replace("$", "").replace("\\", "")
    return re.sub(r"\s+", " ", s).strip()


def failed_list():
    return list(_FAILED)


def _selftest():
    import copy
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt

    cases = [
        r"\frac{\partial \Delta D_t}{\partial p_t}=\bar\theta\kappa S_{t-1}>0",
        r"X'_{\bar e}\eta-\kappa_1\alpha-\kappa_2\eta",
        r"\frac{\partial Net}{\partial X}>0 \quad \text{人民币贬值}",
        r"S_t = S_{t-1}(1+g_t)",
    ]
    doc = Document()
    ok = 0
    for i, c in enumerate(cases, 1):
        p = doc.add_paragraph()
        add_math_run(p, c)
        xml = p._p.xml
        n = xml.count("<m:oMath")
        has_struct = any(t in xml for t in STRUCT_TAGS)
        toks = re.findall(r"<m:t[^>]*>([^<]*)</m:t>", xml)
        gloss = "（人民币贬值）" in p.text
        flag = "OK " if n else "FAIL"
        ok += bool(n)
        print(f"  [{flag}] 式{i}: oMath={n} 结构={has_struct} 词元={toks[:8]}" + (f" GLOSS={gloss}" if gloss else ""))
    print(f"\n自检通过 {ok}/{len(cases)}；公式转换失败 {len(_FAILED)} 条")
    for k, e in _FAILED:
        print("   失败:", k, "|", e)
    return 0 if ok == len(cases) else 1


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        sys.exit(_selftest())
    src = " ".join(a for a in sys.argv[1:] if not a.startswith("-")) or r"\frac{a}{b}"
    cleaned, glosses = clean_tex(src)
    el = omml_for(cleaned)
    if el is None:
        print("FAIL: pandoc 未产出 oMath")
        sys.exit(2)
    import lxml.etree as ET
    xml = ET.tostring(el).decode()
    print("清洗后:", cleaned)
    if glosses:
        print("摘出中文注解:", glosses)
    print("OMML 词元:", re.findall(r"<m:t[^>]*>([^<]*)</m:t>", xml)[:20])
    print("结构标签:", [t for t in STRUCT_TAGS if "<" + t in xml])
